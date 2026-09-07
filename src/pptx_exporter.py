from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd
from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from src.localization import (
    glossary_text,
    localize_factsheet,
    risk_label,
    t,
    translate_period_label,
    translate_value,
)


A4_W = 8.2677
A4_H = 11.6929
BLACK = RGBColor(0, 0, 0)
DARK_GRAY = RGBColor(74, 74, 74)
MID_GRAY = RGBColor(119, 119, 119)
LIGHT_GRAY = RGBColor(232, 232, 232)
PALE_GRAY = RGBColor(244, 244, 244)
WHITE = RGBColor(255, 255, 255)

PERCENT_RISK = {"annualized_volatility", "maximum_drawdown", "tracking_error"}


def _pct(value: float | None) -> str:
    if value is None or pd.isna(value):
        return "-"
    return f"{float(value):.2%}"


def _ratio(value: float | None) -> str:
    if value is None or pd.isna(value):
        return "-"
    return f"{float(value):.2f}"


def _currency(value: float | None, currency: str) -> str:
    if value is None or pd.isna(value):
        return "-"
    if abs(value) >= 1_000_000_000:
        return f"{currency} {value / 1_000_000_000:,.2f}bn"
    if abs(value) >= 1_000_000:
        return f"{currency} {value / 1_000_000:,.2f}m"
    return f"{currency} {value:,.2f}"


def _font(language: str) -> str:
    return "Microsoft YaHei" if language == "zh" else "Arial"


def _set_run_font(run, language: str, size: float, bold=False, color=BLACK):
    run.font.name = _font(language)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    language,
    size=8,
    bold=False,
    color=BLACK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    _set_run_font(run, language, size, bold=bold, color=color)
    return box


def _add_line(slide, x1, y1, x2, y2, width=0.55, color=DARK_GRAY):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def _add_section_title(slide, x, y, w, text, language):
    _add_text(slide, x, y, w, 0.22, text, language, size=8.2, bold=True)


def _add_header(slide, fund, report_date, language, logo_bytes=None, page=1):
    left, right, top = 0.39, A4_W - 0.39, 0.39
    if page == 1:
        ticker_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.70), Inches(0.70))
        ticker_box.fill.solid()
        ticker_box.fill.fore_color.rgb = RGBColor(238, 238, 238)
        ticker_box.line.fill.background()
        tf = ticker_box.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(fund.get("Ticker") or "FUND")[:8]
        _set_run_font(run, language, 9.5, bold=True)
        title_x = left + 0.93
    else:
        title_x = left

    _add_text(
        slide,
        title_x,
        top + 0.08,
        5.6,
        0.38,
        str(fund.get("Fund_Name", "Fund")),
        language,
        size=15 if page == 1 else 12.5,
        bold=True,
    )

    if logo_bytes and page == 1:
        stream = BytesIO(logo_bytes)
        with PILImage.open(BytesIO(logo_bytes)) as im:
            iw, ih = im.size
        max_w, max_h = 1.45, 0.60
        scale = min(max_w / iw, max_h / ih)
        pic_w, pic_h = iw * scale, ih * scale
        slide.shapes.add_picture(
            stream,
            Inches(right - pic_w),
            Inches(top + (max_h - pic_h) / 2),
            width=Inches(pic_w),
            height=Inches(pic_h),
        )

    if page == 1:
        if language == "zh":
            date_str = f"{report_date.year}年{report_date.month}月{report_date.day}日"
        else:
            date_str = report_date.strftime("%B %d, %Y")
        _add_text(slide, left, top + 0.88, 4.3, 0.25, t("fact_sheet_as_of", language, date=date_str), language, size=8.6, bold=True)
        line_y = top + 1.13
    else:
        line_y = top + 0.79

    _add_line(slide, left, line_y, right, line_y, width=0.6)
    return line_y + 0.14


def _add_footer(slide, page, language):
    _add_text(slide, 0.39, 11.34, 4.7, 0.14, t("footer", language), language, size=5.5, color=RGBColor(85, 85, 85))
    _add_text(slide, 7.0, 11.34, 0.88, 0.14, t("page", language, page=page), language, size=5.5, color=RGBColor(85, 85, 85), align=PP_ALIGN.RIGHT)


def _set_cell(cell, text, language, size=6.2, bold=False, align=PP_ALIGN.LEFT, fill=None):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
    cell.margin_left = Pt(2)
    cell.margin_right = Pt(2)
    cell.margin_top = Pt(1.2)
    cell.margin_bottom = Pt(1.2)
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    _set_run_font(run, language, size, bold=bold)


def _add_performance_table(slide, df, language, x, y, w, h):
    columns = list(df.columns)
    rows = []
    header = [translate_period_label(str(c), language) for c in columns]
    if columns and columns[0] == "Series":
        header[0] = t("series", language)
    rows.append(header)
    for _, row in df.iterrows():
        label = str(row[columns[0]])
        if label == "Fund":
            label = t("fund", language)
        elif label == "Benchmark":
            label = t("benchmark", language)
        rows.append([label] + [_pct(row[c]) for c in columns[1:]])

    shape = slide.shapes.add_table(len(rows), len(columns), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    first_w = w * 0.23
    table.columns[0].width = Inches(first_w)
    for c in range(1, len(columns)):
        table.columns[c].width = Inches((w - first_w) / max(1, len(columns) - 1))

    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            fill = None
            bold = r == 0 or c == 0
            if c == 0 and r > 0:
                if value == t("fund", language):
                    fill = LIGHT_GRAY
                elif value == t("benchmark", language):
                    fill = PALE_GRAY
            _set_cell(table.cell(r, c), value, language, size=5.7 if language == "zh" else 5.9, bold=bold, align=PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT, fill=fill)
    return shape


def _add_kv_rows(slide, rows, language, x, y, w, line_h=0.24):
    for i, (label, value) in enumerate(rows):
        yy = y + i * line_h
        _add_text(slide, x, yy, w * 0.58, line_h, label, language, size=6.4, bold=True)
        _add_text(slide, x + w * 0.58, yy, w * 0.42, line_h, value, language, size=6.4)
    return y + len(rows) * line_h


def _add_top_holdings(slide, top, language, x, y, w, max_rows=8):
    if top is None:
        return y
    for i, (_, row) in enumerate(top.head(max_rows).iterrows()):
        yy = y + i * 0.21
        _add_text(slide, x, yy, w * 0.74, 0.20, str(row.get("Security_Name", "")), language, size=6.2)
        _add_text(slide, x + w * 0.74, yy, w * 0.26, 0.20, _pct(row.get("Weight")), language, size=6.2, align=PP_ALIGN.RIGHT)
    return y + min(max_rows, len(top)) * 0.21


def _add_growth_chart(slide, growth, language, x, y, w, h):
    working = growth.copy().set_index("Date")
    monthly = working[["Fund_Growth", "Benchmark_Growth"]].resample("ME").last().dropna().reset_index()
    if len(monthly) > 90:
        monthly = monthly.iloc[::2].copy()

    chart_data = ChartData()
    chart_data.categories = [d.strftime("%Y-%m") for d in monthly["Date"]]
    chart_data.add_series(t("fund", language), monthly["Fund_Growth"].tolist())
    chart_data.add_series(t("benchmark", language), monthly["Benchmark_Growth"].tolist())
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE,
        Inches(x), Inches(y), Inches(w), Inches(h),
        chart_data,
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(5)
    chart.value_axis.tick_labels.font.size = Pt(5)
    for series in chart.series:
        series.format.line.width = Pt(1.4)
    return chart


def _risk_period_label(factsheet, language):
    risk = factsheet.get("risk", {})
    if not risk:
        return ""
    result = next(iter(risk.values()))
    if result.start_date is None or result.end_date is None:
        return ""
    years = (result.end_date - result.start_date).days / 365.25
    if 0.8 <= years <= 1.2:
        return "1年" if language == "zh" else "1y"
    if 2.7 <= years <= 3.3:
        return "3年" if language == "zh" else "3y"
    if 4.6 <= years <= 5.4:
        return "5年" if language == "zh" else "5y"
    return ""


def _add_page1(prs, factsheet, language, logo_bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fund = factsheet["fund"]
    report_date = pd.Timestamp(factsheet["reporting_date"])
    body_top = _add_header(slide, fund, report_date, language, logo_bytes, page=1)

    left_x, left_w, gap = 0.42, 4.73, 0.12
    right_x, right_w = left_x + left_w + gap, 2.56
    y = body_top

    # Left column
    _add_section_title(slide, left_x, y, left_w, t("fund_description", language), language)
    desc = str(fund.get("Fund_Description") or "-")
    _add_text(slide, left_x, y + 0.23, left_w, 0.48, desc, language, size=6.5)
    _add_line(slide, left_x, y + 0.74, left_x + left_w, y + 0.74)
    y += 0.84

    perf = factsheet.get("performance", {})
    growth = perf.get("growth_of_10000")
    if growth is not None:
        _add_section_title(slide, left_x, y, left_w, t("growth_title", language), language)
        _add_growth_chart(slide, growth, language, left_x, y + 0.25, left_w, 2.15)
        _add_text(slide, left_x, y + 2.43, left_w, 0.36, t("growth_note", language), language, size=5.4, color=RGBColor(70, 70, 70))
        y += 2.88

    calendar = perf.get("calendar_year_returns")
    if calendar is not None:
        pivot = calendar.pivot(index="Series", columns="Period", values="Return").reset_index()
        _add_section_title(slide, left_x, y, left_w, t("calendar_performance", language), language)
        _add_performance_table(slide, pivot, language, left_x, y + 0.25, left_w, 0.77)
        y += 1.12

    standard = perf.get("standard_performance_table")
    if standard is not None:
        _add_section_title(slide, left_x, y, left_w, t("annualized_performance", language), language)
        _add_performance_table(slide, standard, language, left_x, y + 0.25, left_w, 0.77)
        y += 1.12

    custom = perf.get("custom_net_performance")
    if custom is not None:
        period = ""
        if custom.start_date is not None and custom.end_date is not None:
            period = f"{custom.start_date.strftime('%b %Y')} - {custom.end_date.strftime('%b %Y')}" if language == "en" else f"{custom.start_date.year}.{custom.start_date.month}-{custom.end_date.year}.{custom.end_date.month}"
        rows = [[t("net_performance", language), period, _pct(custom.value)]]
        shape = slide.shapes.add_table(1, 3, Inches(left_x), Inches(y), Inches(left_w), Inches(0.36))
        tbl = shape.table
        widths = [1.5, 2.05, left_w - 3.55]
        for c, ww in enumerate(widths):
            tbl.columns[c].width = Inches(ww)
            _set_cell(tbl.cell(0, c), rows[0][c], language, size=5.8, bold=c == 0, align=PP_ALIGN.CENTER if c == 2 else PP_ALIGN.LEFT, fill=LIGHT_GRAY if c == 0 else None)
        y += 0.48

    _add_text(slide, left_x, y, left_w, 0.55, t("past_performance_note", language), language, size=5.4, color=RGBColor(70, 70, 70))
    _add_line(slide, left_x, y + 0.58, left_x + left_w, y + 0.58)

    # Right column
    ry = body_top
    latest = factsheet.get("latest_nav") or {}
    _add_section_title(slide, right_x, ry, right_w, t("key_facts", language), language)
    launch = "-"
    if pd.notna(fund.get("Launch_Date")):
        d = pd.Timestamp(fund.get("Launch_Date"))
        launch = f"{d.year}/{d.month}/{d.day}" if language == "zh" else d.strftime("%m/%d/%Y")
    key_rows = [
        (t("asset_class", language) + " :", translate_value(fund.get("Asset_Class"), language)),
        (t("benchmark", language) + " :", str(fund.get("Benchmark_Name", "-"))),
        (t("launch_date", language) + " :", launch),
        (t("currency", language) + " :", str(fund.get("Currency", "-"))),
        (t("ticker", language) + " :", str(fund.get("Ticker", "-"))),
        (t("net_assets", language) + " :", _currency(latest.get("AUM"), str(fund.get("Currency", "")))),
    ]
    ry = _add_kv_rows(slide, key_rows, language, right_x, ry + 0.27, right_w, line_h=0.31)
    _add_line(slide, right_x, ry + 0.03, right_x + right_w, ry + 0.03)
    ry += 0.14

    _add_section_title(slide, right_x, ry, right_w, t("fees", language), language)
    ry = _add_kv_rows(slide, [(t("management_fee", language), _pct(fund.get("Management_Fee")))], language, right_x, ry + 0.26, right_w)
    _add_line(slide, right_x, ry + 0.04, right_x + right_w, ry + 0.04)
    ry += 0.15

    _add_section_title(slide, right_x, ry, right_w, t("fund_characteristics", language), language)
    period = _risk_period_label(factsheet, language)
    chars = []
    for metric_id, result in factsheet.get("risk", {}).items():
        label = risk_label(metric_id, language)
        if period:
            label += f" ({period})"
        val = _pct(result.value) if metric_id in PERCENT_RISK else _ratio(result.value)
        chars.append((label + " :", val))
    portfolio = factsheet.get("portfolio", {})
    if "number_of_holdings" in portfolio:
        chars.append((t("number_holdings", language) + " :", str(portfolio["number_of_holdings"])))
    if "cash_weight" in portfolio:
        chars.append((t("cash_weight", language) + " :", _pct(portfolio["cash_weight"])))
    if "top_10_concentration" in portfolio:
        chars.append((t("top10_concentration", language) + " :", _pct(portfolio["top_10_concentration"])))
    ry = _add_kv_rows(slide, chars, language, right_x, ry + 0.28, right_w, line_h=0.29)
    _add_line(slide, right_x, ry + 0.03, right_x + right_w, ry + 0.03)
    ry += 0.14

    _add_section_title(slide, right_x, ry, right_w, t("top_holdings", language), language)
    ry = _add_top_holdings(slide, portfolio.get("top_holdings"), language, right_x, ry + 0.25, right_w, max_rows=8)
    _add_text(slide, right_x, ry + 0.02, right_w, 0.26, t("holdings_change", language), language, size=5.3, color=RGBColor(70, 70, 70))
    _add_line(slide, right_x, ry + 0.31, right_x + right_w, ry + 0.31)

    _add_footer(slide, 1, language)
    return slide


def _add_allocation_table(slide, df, category, language, x, y, w, title):
    _add_section_title(slide, x, y, w, title, language)
    if df is None or df.empty:
        return y + 0.3
    rows = df[[category, "Weight"]].head(10)
    shape = slide.shapes.add_table(len(rows), 2, Inches(x), Inches(y + 0.27), Inches(w), Inches(min(2.1, len(rows) * 0.22)))
    tbl = shape.table
    tbl.columns[0].width = Inches(w * 0.72)
    tbl.columns[1].width = Inches(w * 0.28)
    for r, (_, row) in enumerate(rows.iterrows()):
        _set_cell(tbl.cell(r, 0), str(row[category]), language, size=5.9)
        _set_cell(tbl.cell(r, 1), _pct(row["Weight"]), language, size=5.9, align=PP_ALIGN.RIGHT)
    return y + 0.27 + min(2.1, len(rows) * 0.22)


def _add_page2(prs, factsheet, language, logo_bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fund = factsheet["fund"]
    report_date = pd.Timestamp(factsheet["reporting_date"])
    body_top = _add_header(slide, fund, report_date, language, logo_bytes, page=2)
    left_x, col_w, gap = 0.42, 3.55, 0.25
    right_x = left_x + col_w + gap

    entries = []
    for metric_id in factsheet.get("risk", {}):
        text = glossary_text(metric_id, language)
        if text:
            entries.append((risk_label(metric_id, language), text))
    entries += [
        (t("net_assets", language), glossary_text("net_assets", language)),
        (t("number_holdings", language), glossary_text("number_holdings", language)),
    ]
    midpoint = (len(entries) + 1) // 2
    left_entries, right_entries = entries[:midpoint], entries[midpoint:]

    _add_section_title(slide, left_x, body_top, col_w, t("glossary", language), language)
    y1 = body_top + 0.28
    for title, text in left_entries:
        _add_text(slide, left_x, y1, col_w, 0.53, f"{title}: {text}", language, size=5.9)
        y1 += 0.58
    y2 = body_top + 0.28
    for title, text in right_entries:
        _add_text(slide, right_x, y2, col_w, 0.53, f"{title}: {text}", language, size=5.9)
        y2 += 0.58

    glossary_bottom = max(y1, y2) + 0.05
    _add_line(slide, left_x, glossary_bottom, 7.85, glossary_bottom, width=0.5, color=MID_GRAY)

    portfolio = factsheet.get("portfolio", {})
    alloc_y = glossary_bottom + 0.14
    left_bottom = alloc_y
    if "sector_allocation" in portfolio:
        left_bottom = _add_allocation_table(slide, portfolio["sector_allocation"], "Sector", language, left_x, alloc_y, col_w, t("sector_allocation", language))
    if "asset_class_allocation" in portfolio:
        left_bottom = _add_allocation_table(slide, portfolio["asset_class_allocation"], "Asset_Class", language, left_x, left_bottom + 0.14, col_w, t("asset_class_allocation", language))
    right_bottom = alloc_y
    if "country_allocation" in portfolio:
        right_bottom = _add_allocation_table(slide, portfolio["country_allocation"], "Country", language, right_x, alloc_y, col_w, t("country_allocation", language))

    info_y = max(left_bottom, right_bottom) + 0.20
    _add_line(slide, left_x, info_y, 7.85, info_y, width=0.5, color=MID_GRAY)
    info_y += 0.14
    _add_section_title(slide, left_x, info_y, 7.43, t("important_info", language), language)
    info_y += 0.27
    source = fund.get("Source")
    lines = []
    if source and not pd.isna(source):
        lines.append(f"{t('source', language)}: {source}.")
    basis = []
    if fund.get("Performance_Basis") and not pd.isna(fund.get("Performance_Basis")):
        basis.append(f"{t('performance_basis', language)}: {fund.get('Performance_Basis')}")
    if fund.get("Benchmark_Return_Type") and not pd.isna(fund.get("Benchmark_Return_Type")):
        basis.append(f"{t('benchmark_basis', language)}: {fund.get('Benchmark_Return_Type')}")
    if basis:
        lines.append(". ".join(basis) + ".")
    lines.append(t("demo_legal_bold", language))
    lines.append(t("demo_legal", language))
    legal_text = "\n\n".join(lines)
    _add_text(slide, left_x, info_y, 7.43, 1.80, legal_text, language, size=5.6)
    _add_footer(slide, 2, language)
    return slide


def generate_factsheet_pptx(
    factsheet: dict[str, Any],
    logo_bytes: bytes | None = None,
    language: str = "en",
    openrouter_api_key: str | None = None,
    openrouter_model: str | None = None,
) -> bytes:
    """Generate an A4 portrait PowerPoint where text, tables, lines and charts remain editable."""
    language = "zh" if language == "zh" else "en"
    factsheet = localize_factsheet(
        factsheet,
        language=language,
        api_key=openrouter_api_key,
        model=openrouter_model,
    )
    prs = Presentation()
    prs.slide_width = Inches(A4_W)
    prs.slide_height = Inches(A4_H)
    _add_page1(prs, factsheet, language, logo_bytes)
    _add_page2(prs, factsheet, language, logo_bytes)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
